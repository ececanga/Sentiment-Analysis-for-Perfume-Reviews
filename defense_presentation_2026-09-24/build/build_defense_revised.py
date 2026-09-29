from pathlib import Path
import json, csv, hashlib, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.chart.data import CategoryChartData
from pptx.oxml.xmlchemy import OxmlElement
from PIL import ImageFont

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'defense_presentation_2026-09-24'
PPTX=OUT/'Ece_CANGA_Thesis_Defense_Revised.pptx'
C={'navy':'173B4D','teal':'14818B','muted':'5F727C','light':'EAF2F3','bg':'F8FAFA','white':'FFFFFF','line':'DCE5E8','blue':'88A2AF','orange':'CC7520','paleorange':'FBF1E5'}
prs=Presentation(); prs.slide_width= Inches(13.333333);prs.slide_height= Inches(7.5)
prs.core_properties.title='Emotions in Perfume Reviews — Thesis Defense'
prs.core_properties.subject='Advanced Sentiment Analysis Framework for Olfactory Perception Mining'
prs.core_properties.author='Ece ÇANGA'
prs.core_properties.keywords='thesis defense, perfume reviews, multi-label emotion classification'
prs.core_properties.comments='30 main slides and 10 backup slides. Editable text, charts, tables and diagrams.'
manifest=[]
DATA=ROOT/'outputs/thesis_consolidated_results_v1'
with (ROOT/'outputs/source_grouped_feature_ablation_v1/feature_columns_by_variant.csv').open() as f:
 FEATURE_COLUMNS=next(r for r in csv.DictReader(f) if r['variant']=='full_hybrid')['feature_columns'].split(';')
assert len(FEATURE_COLUMNS)==58
with (DATA/'youtube_five_model_pooled.csv').open() as f: NATURAL=list(csv.DictReader(f))
with (DATA/'feature_combination_three_model_pooled.csv').open() as f: ABLATION=list(csv.DictReader(f))


def rgb(c):return RGBColor.from_string(C.get(c,c))
def rect(s,x,y,w,h,fill='white',line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
 # Suppress theme-inherited shape shadows.
 sh._element.spPr.append(OxmlElement('a:effectLst'))
 if line:sh.line.color.rgb=rgb(line);sh.line.width=Pt(.7)
 else:sh.line.fill.background()
 if radius:
  try:sh.adjustments[0]=.08
  except:pass
 return sh

def txt(s,x,y,w,h,text,size=22,color='navy',bold=False,align=PP_ALIGN.LEFT,val=MSO_ANCHOR.TOP,margin=0):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
 tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right= Inches(margin);tf.margin_top=tf.margin_bottom= Inches(margin);tf.vertical_anchor=val
 for i,t in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.font.name='Arial';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.alignment=align;p.space_after=Pt(5);p.line_spacing=1.12
 return sh

def base(num,title,section,source='',dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('navy' if dark else 'bg')
 col='white' if dark else 'navy'
 rect(s,.55,.38,.07,.22,'teal');txt(s,.76,.36,11,.3,section.upper(),11,'blue' if dark else 'teal',True)
 # measured title size to keep long headings readable within two lines
 size=32 if len(title)<69 else 30
 txt(s,.65,.94,12,1.13,title,size,col,True)
 rect(s,.65,6.96,12.03,.012,'blue' if dark else 'line')
 if source:txt(s,.65,7.09,10.9,.22,source,9.2,'blue' if dark else 'muted')
 txt(s,11.9,7.06,.75,.26,str(num).zfill(2) if isinstance(num,int) else num,11,'blue' if dark else 'muted',align=PP_ALIGN.RIGHT)
 manifest.append({'slide':len(prs.slides),'label':num,'title':title,'source':source})
 s.notes_slide.notes_text_frame.text=f'{num} — {title}\nSource: {source}'
 return s

def band(s,text,y=6.03,fill='light',size=21,color='navy',h=.67):
 h=max(h,.54);y=min(y,6.82-h)
 font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',size*4)
 needed=font.getlength(text)/4
 size=min(size, max(18,int(size*11.55*72/needed)))
 rect(s,.65,y,12.03,h,fill,radius=True)
 txt(s,.87,y+.07,11.55,h-.1,text,size,color,val=MSO_ANCHOR.MIDDLE)

def card(s,x,y,w,h,heading,body,accent='teal',body_size=21):
 rect(s,x,y,w,h,'white',line='line',radius=True);rect(s,x,y,.055,h,accent)
 head_h=.84 if '\n' in heading or len(heading)>25 and w<4 else .65
 body_y=1.29 if head_h>.7 else 1.02
 if h<2:body_y=.76
 txt(s,x+.23,y+.22,w-.46,head_h,heading,23,accent,True)
 txt(s,x+.23,y+body_y,w-.46,h-body_y-.12,body,body_size)

def pill(s,x,y,w,label,fill='light',size=18):
 rect(s,x,y,w,.5,fill,radius=True);txt(s,x+.07,y+.09,w-.14,.33,label,size,'navy',True,PP_ALIGN.CENTER)

def arrow(s,x1,y1,x2,y2,color='blue'):
 c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));c.line.color.rgb=rgb(color);c.line.width=Pt(1.7)
 ln=c._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln');tail=OxmlElement('a:tailEnd');tail.set('type','triangle');tail.set('w','med');tail.set('len','med');ln.append(tail)
 return c

def node(s,x,y,w,h,label,fill='light',size=20,bold=True):
 rect(s,x,y,w,h,fill,radius=True);txt(s,x+.12,y+.08,w-.24,h-.16,label,size,'white' if fill=='navy' else 'navy',bold,PP_ALIGN.CENTER,MSO_ANCHOR.MIDDLE)

def table(s,headers,rows,x=.65,y=2.1,w=12.03,row_h=.67,widths=None,size=20,highlights=None):
 widths=widths or [w/len(headers)]*len(headers)
 sh=s.shapes.add_table(len(rows)+1,len(headers), Inches(x),Inches(y), Inches(w),Inches(row_h*(len(rows)+1)));tb=sh.table
 for j,cw in enumerate(widths):tb.columns[j].width=Inches(cw)
 for i,row in enumerate([headers]+rows):
  tb.rows[i].height=Inches(row_h)
  for j,v in enumerate(row):
   c=tb.cell(i,j);c.text=str(v);c.margin_left=c.margin_right=Inches(.16);c.margin_top=c.margin_bottom=Inches(.04);c.vertical_anchor=MSO_ANCHOR.MIDDLE
   c.fill.solid();c.fill.fore_color.rgb=rgb('navy' if i==0 else ('white' if i%2 else 'light'))
   for p in c.text_frame.paragraphs:
    p.font.name='Arial';p.font.size=Pt(size if i else max(size-1,15));p.font.color.rgb=rgb('white' if i==0 else 'navy');p.font.bold=(i==0 or bool(highlights and (i-1,j) in highlights));p.space_after=Pt(0);p.line_spacing=1.04
   tcpr=c._tc.get_or_add_tcPr()
   for edge in ['lnL','lnR','lnT','lnB']:
    ln=OxmlElement('a:'+edge);ln.set('w','0');nf=OxmlElement('a:noFill');ln.append(nf);tcpr.append(ln)
 return sh

def chart(s,categories,series,x=.8,y=2.13,w=11.5,h=3.55,max_val=.4,fmt='0.0000',colors=None,highlight=None,horizontal=False):
 data=CategoryChartData();data.categories=categories
 for name,vals in series:data.add_series(name,vals)
 sh=s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(x),Inches(y),Inches(w),Inches(h),data);ch=sh.chart
 ch.chart_style=10;ch.has_legend=len(series)>1
 if ch.has_legend:ch.legend.position=XL_LEGEND_POSITION.BOTTOM;ch.legend.include_in_layout=False;ch.legend.font.name='Arial';ch.legend.font.size=Pt(15)
 ch.has_title=False
 for ax in [ch.category_axis,ch.value_axis]:
  ax.tick_labels.font.name='Arial';ax.tick_labels.font.size=Pt(15);ax.tick_labels.font.color.rgb=rgb('navy');ax.major_tick_mark=XL_TICK_MARK.NONE;ax.minor_tick_mark=XL_TICK_MARK.NONE
  ax.format.line.fill.background()
 va=ch.value_axis;va.minimum_scale=0;va.maximum_scale=max_val;va.major_unit=100 if max_val>1 else .1;va.tick_labels.number_format='0' if max_val>1 else '0.0';va.has_major_gridlines=True;va.major_gridlines.format.line.color.rgb=rgb('line');va.major_gridlines.format.line.width=Pt(.5)
 ch.category_axis.has_major_gridlines=False
 plot=ch.plots[0];plot.gap_width=65 if len(series)==1 else 55;plot.has_data_labels=True;dl=plot.data_labels;dl.position=XL_LABEL_POSITION.OUTSIDE_END;dl.number_format='[$-en-US]'+fmt;dl.font.name='Arial';dl.font.size=Pt(15);dl.font.color.rgb=rgb('navy')
 for k,ser in enumerate(ch.series):
  col=(colors or ['teal','orange'])[k];ser.format.fill.solid();ser.format.fill.fore_color.rgb=rgb(col);ser.format.line.fill.background()
  if highlight is not None:
   for i,pt in enumerate(ser.points):pt.format.fill.solid();pt.format.fill.fore_color.rgb=rgb('teal' if i==highlight else 'blue');pt.format.line.fill.background()
 return sh

# 01
s=base(1,'Emotions in\nPerfume Reviews','Master’s thesis defense',dark=True)
# cover custom larger title, while preserving base header
for sh in s.shapes:
 if sh.has_text_frame and sh.text=='Emotions in\nPerfume Reviews':
  sh.top=Inches(1.03);sh.height=Inches(1.6)
  for p in sh.text_frame.paragraphs:p.font.size=Pt(45)
txt(s,.69,2.96,11.3,1.3,'Advanced Sentiment Analysis Framework for Olfactory Perception Mining:\nUtilizing Machine Learning and Deep Learning Techniques\non Social Media Data Streams',21,'white')
rect(s,.69,4.59,1.0,.04,'teal')
txt(s,.69,4.95,6,.45,'Ece ÇANGA',26,'white',True)
txt(s,.69,5.55,6,.83,'Computer Engineering\nÇankaya University',19,'white')
txt(s,7.5,5.55,5,.9,'Supervisor: Assoc. Prof. Dr. Serdar ARSLAN\nSeptember 2026',17,'white')
# 02
s=base(2,'Why look beyond positive and negative reviews?','Motivation','Nikitina & Furuoka (2023); Nandwani & Verma (2021).')
txt(s,.65,2.03,12,0.7,'Perfume reviews describe sensory impressions and emotional associations.',24)
card(s,.65,3.02,5.88,2.09,'Enjoyment','“This makes me happy.”',body_size=25)
card(s,6.79,3.02,5.88,2.09,'Anticipation','“I cannot wait to wear this.”',body_size=25)
txt(s,.8,5.29,11,.34,'Two positive reviews; different emotional meanings. Illustrative examples.',16,'muted')
band(s,'Emotion classification offers a more detailed description of review language.',size=22)
#03
s=base(3,'One review can express several emotions','The task')
rect(s,.65,2.03,12.03,1.4,'white',line='line',radius=True)
txt(s,.96,2.24,11.4,.95,'“This perfume makes me happy, and I cannot wait to wear it.”',27,bold=True)
txt(s,.73,3.51,3,.3,'Illustrative example',14,'muted');pill(s,7.6,3.52,1.75,'Joy');txt(s,9.5,3.56,.3,.4,'+',23);pill(s,9.95,3.52,2.5,'Anticipation')
txt(s,.65,4.28,12,.76,'The task allows multiple emotion labels for the same perfume, based on its review text.',23)
for i,(lab,w) in enumerate([('Joy',1.2),('Trust',1.3),('Anticipation',2.25),('Surprise',1.65),('Sadness',1.6),('Fear',1.2),('Disgust',1.55)]):
 xx=.65+sum([1.2,1.3,2.25,1.65,1.6,1.2,1.55][:i])+.19*i;pill(s,xx,5.29,w,lab,size=17)
band(s,'Target: emotion expressed in review language. Wearer responses are not directly measured.',size=20)
#04
s=base(4,'Existing methods provide different kinds of evidence','Background','Mohammad & Turney (2013); Mohammad (2018); Demszky et al. (2020); Devlin et al. (2019).')
table(s,['Approach','What it contributes'],[['Emotion dictionaries','Associations between words and emotions'],['Affect dimensions','Pleasantness, activation and control'],['Contextual language models','Patterns that depend on surrounding text']],row_h=.91,widths=[4.8,7.23],size=23)
band(s,'This thesis evaluates these approaches within a shared perfume-review benchmark.',size=21)
#05
s=base(5,'Three research questions','Research aim','Thesis §1.4.')
for y,tag,title,body in [(2.03,'RQ1','Representation','How can perfume-review transcripts be represented as structured emotion features?'),(3.55,'RQ2','Model comparison','How do machine learning and deep learning methods compare in identifying multiple emotions?'),(5.07,'RQ3','Synthetic augmentation','How does weighted augmentation with generated reviews affect prediction on held-out YouTube perfumes?')]:
 node(s,.65,y,1.12,1.2,tag,'navy',23);txt(s,2.04,y+.02,10.35,.38,title,23,'teal',True);txt(s,2.04,y+.49,10.35,.76,body,22)
#06
s=base(6,'What this thesis contributes','Contribution','Thesis §§1.5 and 6.3.')
card(s,.65,2.17,3.84,3.6,'A perfume-review\ndataset','English YouTube transcript records linked to 784 perfumes.',body_size=22)
card(s,4.74,2.17,3.84,3.6,'A reproducible\nworkflow','A 58-feature representation alongside a direct text modelling route.',body_size=22)
card(s,8.83,2.17,3.84,3.6,'A controlled\ncomparison','Five models, feature combinations and weighted synthetic augmentation.',body_size=22)
band(s,'Shared evaluation conditions connect the dataset, representations and experiments.',size=21)
#07
s=base(7,'From review transcripts to perfume-level predictions','Data','Thesis §3.2.1; Table 1.')
for x,value,label in [(.65,'784','perfumes'),(4.78,'4,881','usable transcript records'),(8.91,'1–10','records per perfume; median 7')]:
 txt(s,x,1.95,3.7,.8,value,44,'teal',True);txt(s,x,2.83,3.7,.72,label,19)
txt(s,.65,3.78,12,.75,'Perfume-specific YouTube searches supplied available English transcripts.\nRetained transcripts contain at least 20 words.',21)
node(s,.65,4.84,3.2,.84,'Transcript records');arrow(s,3.93,5.25,4.54,5.25);node(s,4.62,4.84,3.45,.84,'Grouped by perfume');arrow(s,8.17,5.25,8.78,5.25);node(s,8.87,4.84,3.8,.84,'Emotion predictions\nfor each perfume',size=19)
band(s,'Perfumes can share source material, so the records are not necessarily independent.',size=21)
#08
s=base(8,'The reference labels follow a fixed annotation policy','Reference labels','Thesis §3.4; Niu et al. (2024), methodological context.')
card(s,.65,2.13,5.85,3.45,'Gemini 3.1 Pro','Assigned emotion labels from supplied review excerpts.\n\nFixed instructions and a consistent label set.',body_size=23)
card(s,6.83,2.13,5.85,3.45,'What the model saw','One or two excerpts per perfume.\n\nMaximum 360 characters per excerpt.',body_size=23)
band(s,'Evaluation measures agreement with these model-generated reference labels.',y=5.88,size=21)
txt(s,.86,6.63,11.3,.26,'Independent human validation was not conducted.',16,'muted')
#09

s=base(9,'Which model does what?','Roles in the workflow','Thesis §§3.2.2 and 3.4–3.6. Evaluation uses the Gemini reference-label policy.')
rows=[
 ['Hartmann (DistilRoBERTa) · Savani (DistilBERT)\nGoEmotions (BERT) · Cardiff (RoBERTa)','Extract 45 numerical emotion scores\nthat enter the 58-feature representation'],
 ['Gemini 3.1 Pro','Assigns reference emotion labels\nfrom the supplied review text'],
 ['Logistic Regression · XGBoost · Random Forest\nBERT · BiLSTM with Attention','Five supervised classifiers:\nlearn to predict the reference labels'],
 ['GPT-5','Generates additional review text\nfor the augmentation experiment']]
t=table(s,['Named models','Role in this thesis'],rows,y=2.0,widths=[6.5,5.53],row_h=.78,size=20).table
for row,h in zip(t.rows,[.52,1.0,.79,1.0,.79]):row.height=Inches(h)
band(s,'Evaluation compares classifier predictions with Gemini labels using F1, AP and ROC-AUC.',y=6.32,size=20,h=.5)
#10

s=base(10,'Two input routes, five named classifiers','Method overview','Simplified from thesis Figure 1; §§3.5–3.7.')
node(s,.65,3.11,1.95,1.08,'Review\ntranscripts','navy',21)
arrow(s,2.68,3.38,3.15,2.77);arrow(s,2.68,3.94,3.15,4.63)
node(s,3.24,2.22,2.24,1.04,'58 features\nper perfume',size=21)
node(s,3.24,4.17,2.24,1.04,'Transcript\ntext',size=21)
arrow(s,5.56,2.75,6.04,2.75);arrow(s,5.56,4.68,6.04,4.68)
node(s,6.13,2.02,3.61,1.46,'Logistic Regression\nXGBoost\nRandom Forest',size=20)
node(s,6.13,4.05,3.61,1.29,'BERT\nBiLSTM with Attention',size=20)
arrow(s,9.83,2.75,10.2,3.41);arrow(s,9.83,4.68,10.2,3.95)
node(s,10.3,3.08,2.37,1.19,'Each model predicts\nseven emotions','navy',19)
txt(s,9.98,4.58,2.84,1.0,'Compare predictions\nwith held-out\nGemini labels',18,'teal',True,PP_ALIGN.CENTER)
band(s,'All five models face the same held-out perfumes and reference labels.',size=22)
#11

s=base(11,'The complete 58-feature representation','All input features','Thesis §3.5; Table 5; saved 58-column feature schema. Each source–name pair is a separate feature.')
family_titles={'NRC':'NRC word counts','VAD':'VAD dimensions','Hartmann':'Hartmann scores','Savani':'Savani scores','GoEmotions':'GoEmotions scores','Cardiff':'Cardiff scores'}
rows=[]
for family in ['NRC','VAD','Hartmann','Savani','GoEmotions','Cardiff']:
 names=[f.split('_',1)[1] for f in FEATURE_COLUMNS if f.startswith(family+'_')]
 rows.append([family_titles[family],str(len(names)),', '.join(names)])
t=table(s,['Source','Count','Every feature in the representation'],rows,y=1.84,widths=[2.82,1.08,8.13],row_h=.62,size=18).table
for row,h in zip(t.rows,[.45,.61,.45,.5,.5,1.34,.45]):row.height=Inches(h)
txt(s,.7,6.63,12,.29,'58 inputs: 10 word counts + 3 affect dimensions + 45 model scores. Prediction targets: 7 emotions.',16,'teal',True)
#12
s=base(12,'Five models, grouped by their inputs','Model comparison','Thesis §3.6; Table 6.')
card(s,.65,2.1,5.85,3.17,'58 structured features','Logistic Regression\nXGBoost\nRandom Forest',body_size=26)
card(s,6.83,2.1,5.85,3.17,'Review text','BERT\nBiLSTM with Attention',body_size=26)
txt(s,.84,5.45,5.45,.63,'Learn from the 58-value representation.',20)
txt(s,7.02,5.45,5.42,.63,'Learn from transcripts; combine predictions at perfume level.',20)
band(s,'The comparison covers complete approaches, including their different inputs.',y=6.21,size=20,h=.53)
#13
s=base(13,'Performance is measured on held-out perfumes','Evaluation','Thesis §§3.7 and 4.1.1.')
node(s,.65,2.19,3.18,1.05,'Training inputs\nand reference labels',size=21);node(s,.65,3.79,3.18,1.05,'Validation data',size=21)
txt(s,4.14,2.48,2.6,.58,'Learn parameters',18,'muted',align=PP_ALIGN.CENTER);txt(s,4.14,4.43,2.6,.65,'Guide training\ndecisions',18,'muted',align=PP_ALIGN.CENTER)
arrow(s,3.91,2.72,7.0,3.18);arrow(s,3.91,4.31,7.0,3.66);node(s,7.09,2.94,2.29,1.0,'Trained model','navy',21);arrow(s,9.47,3.43,10.0,3.43);node(s,10.1,2.94,2.57,1.0,'Predict test\nperfumes',size=21)
txt(s,.69,5.23,12,.65,'Five outer folds share the same splits. Each perfume is held out for its prediction.',23)
band(s,'Test data do not guide training. Emotion labels use a fixed 0.30 probability cutoff.',size=21)
#14
s=base(14,'Shared sources stay together','Evaluation integrity','Thesis §3.7; Table 8.')
node(s,.65,2.22,2.63,.86,'Perfume A');node(s,.65,4.38,2.63,.86,'Perfume B');node(s,4.33,3.29,3.65,.9,'Shared video or\nidentical full transcript',size=20);arrow(s,3.35,2.64,4.25,3.55);arrow(s,3.35,4.78,4.25,3.9);arrow(s,8.06,3.73,8.85,3.73);node(s,8.94,3.05,3.73,1.34,'Same connected group\nSame evaluation fold','navy',22)
txt(s,4.34,2.14,5.8,.8,'359 source groups',33,'teal',True)
txt(s,.7,5.56,12,.71,'Shared videos and identical transcripts cannot cross training and test data.',22)
txt(s,.7,6.39,11.95,.43,'Groups differ in size; supplementary analyses give each source group equal weight.',18,'muted')
#15

s=base(15,'Three ways to evaluate an emotion prediction','Reading the results','Thesis §4.1.3; Saito & Rehmsmeier (2015). Reference labels are assigned by Gemini.')
txt(s,.68,1.97,12,.53,'Each emotion receives a score. A cutoff converts scores into yes/no labels.',22)
rows=[
 ['Macro F1','Final yes/no labels','At cutoff 0.30, how well do decisions balance missed labels and false alarms?'],
 ['Macro Average\nPrecision (AP)','Ranked scores','Start with the highest scores. AP rewards finding labelled perfumes while limiting false alarms.'],
 ['Macro ROC-AUC','Ranked scores','How often does a labelled perfume score higher than an unlabelled perfume?']]
t=table(s,['Metric','Evaluates','Plain-language meaning'],rows,y=2.73,widths=[3.1,2.4,6.53],row_h=.86,size=20).table
t.rows[0].height=Inches(.53)
for r in list(t.rows)[1:]:r.height=Inches(.96)
txt(s,.72,6.25,11.9,.48,'AP is the primary ranking measure. “Macro” gives each of the seven emotions equal weight.',19,'teal',True)
#16

s=base(16,'A real example of the 58-feature route','Worked example','Thesis §3.9; Table 7. This is one illustrative prediction, not a performance estimate.')
txt(s,.65,1.95,12,.45,'Vetiver Nocturne Absolu Extrait de Parfum',25,'teal',True)
rect(s,.65,2.66,5.03,2.31,'white',line='line',radius=True)
txt(s,.92,2.89,4.49,.86,'“slightly sweet medicinal herbs”',24,bold=True)
txt(s,.92,3.99,4.48,.86,'“Sweet” contributes to dictionary emotion counts.',21)
node(s,.65,5.22,1.41,.71,'10 transcripts',size=15);arrow(s,2.12,5.57,2.45,5.57);node(s,2.51,5.22,1.34,.71,'58 features',size=16);arrow(s,3.91,5.57,4.24,5.57);node(s,4.31,5.22,1.36,.71,'Logistic\nRegression','navy',14)
table(s,['Emotion','Score','At 0.30'],[['Joy','0.674','Present'],['Trust','0.604','Present'],['Anticipation','0.603','Present']],x=6.07,y=2.67,w=6.6,widths=[2.43,1.55,2.62],row_h=.65,size=21)
txt(s,6.1,5.48,6.51,.7,'The three predicted labels match the reference.\nAll five models are compared next.',20)
band(s,'Logistic Regression, XGBoost and Random Forest all use this 58-feature representation.',y=6.3,size=20,h=.51)
#17

s=base(17,'Compare all five models across the two input routes','Natural-data results','Thesis Tables 9 and 10. Pooled out-of-fold Macro AP; the same 784 held-out perfumes.')
rect(s,.78,1.99,.17,.17,'teal');txt(s,1.05,1.94,7,.65,'58 features: Logistic Regression, XGBoost, Random Forest',18,'teal',True)
rect(s,8.12,1.99,.17,.17,'navy');txt(s,8.41,1.94,4.26,.61,'Transcript text: BERT, BiLSTM',18,'navy',True)
ch=chart(s,['Logistic\nRegression','XGBoost','Random\nForest','BERT','BiLSTM +\nAttention'],[('Macro AP',[float(r['macro_average_precision']) for r in NATURAL])],y=2.65,h=3.16).chart
ch.value_axis.has_title=True
ch.value_axis.axis_title.text_frame.text='Macro AP'
for pp in ch.value_axis.axis_title.text_frame.paragraphs:
 pp.font.name='Arial';pp.font.size=Pt(16);pp.font.color.rgb=rgb('navy')
for i,pt in enumerate(ch.series[0].points):
 pt.format.fill.solid();pt.format.fill.fore_color.rgb=rgb('teal' if i<3 else 'navy')
band(s,'Observed Macro AP: Random Forest leads the 58-feature route; BERT leads overall.',size=21)
#18

s=base(18,'All five models across the three metrics','Connecting inputs and results','Thesis Tables 9 and 10; §4.4.2. Pooled out-of-fold results. Bold values are the highest in each metric.')
rows=[]
for i,r in enumerate(NATURAL):
 rows.append([r['model_name'],'58 features' if i<3 else 'Transcript text',*[f"{float(r[k]):.4f}" for k in ['macro_f1','macro_average_precision','macro_roc_auc']]])
table(s,['Model','Input','Macro F1','Macro AP','Macro ROC-AUC'],rows,y=2.07,widths=[3.51,2.4,1.82,1.82,2.48],row_h=.63,size=20,highlights={(3,2),(3,3),(1,4)})
txt(s,.73,6.0,11.95,.45,'BERT leads observed F1 and AP; XGBoost, using all 58 features, leads observed ROC-AUC.',19,'teal',True)
txt(s,.73,6.56,11.95,.31,'Equal-source-group F1 comparisons do not clearly separate the models. Detailed uncertainty: A10.',17,'muted')
#19
s=base(19,'Emotion frequencies shape the evidence','Class imbalance','Thesis Table 14; §4.6.1. Counts are perfume-level reference positives.')
chart(s,['Joy','Trust','Anticipation','Surprise','Disgust','Sadness','Fear'],[('Reference positives',[622,476,366,44,19,11,3])],x=.78,y=2.03,w=8.04,h=3.73,max_val=700,fmt='0',highlight=0)
rect(s,9.16,2.13,3.51,3.8,'light',radius=True);txt(s,9.42,2.42,3,.8,'3',42,'teal',True);txt(s,9.42,3.22,3,1.1,'positive cases\nfor Fear',23,bold=True);txt(s,9.42,4.62,3,1.07,'BERT’s F1 of 0.5000 reflects one correct detection.',20)
band(s,'Joy, Trust and Anticipation dominate; rare-emotion scores depend on very few cases.',size=20)
#20

s=base(20,'Read F1, AP and ROC-AUC together','Interpreting the same predictions','Thesis Tables 10 and 13–15. BERT; natural-only training. Macro scores average all seven emotions.')
txt(s,.69,1.96,12,.48,'BERT: Macro F1 0.4029  ·  Macro AP 0.3590  ·  Macro ROC-AUC 0.5533',23,'teal',True)
txt(s,.68,2.62,6.8,.4,'Final labels at cutoff 0.30',22,bold=True)
table(s,['Emotion','Reference\npositives','Predicted\npositives','F1'],[['Joy','622','784','0.8848'],['Trust','476','784','0.7556'],['Anticipation','366','784','0.6365']],x=.65,y=3.12,w=6.6,widths=[2.0,1.63,1.63,1.34],row_h=.59,size=18)
txt(s,.72,5.53,6.4,.64,'High F1 for common labels can occur when every perfume receives those labels.',19)
rect(s,7.63,2.68,5.04,3.63,'white',line='line',radius=True)
txt(s,7.86,2.91,4.58,.38,'Score rankings, before the cutoff',21,'teal',True)
txt(s,7.86,3.59,4.58,1.04,'AP: rewards finding labelled perfumes while limiting false alarms.',21)
txt(s,7.86,4.79,4.58,.92,'ROC-AUC: compares scores of labelled and unlabelled cases.',21)
band(s,'Changing the cutoff changes F1. AP and ROC-AUC stay fixed when the scores stay fixed.',y=6.39,size=19,h=.43)
#21

s=base(21,'Which feature groups help each classifier?','Feature-combination experiment','Thesis §4.5.1; Table 12. Highest observed values among seven fixed-setting comparisons; not retuned optima.')
txt(s,.68,1.97,12,.43,'Words: 10 features  ·  Affect dimensions: 3  ·  Transformer scores: 45',21,'teal',True)
txt(s,.68,2.49,12,.46,'Test each group alone, each pair, and all three: seven combinations per classifier.',21)
rows=[]
short={'full_hybrid':'All 58 features','nrc_only':'Words (10)','vad_only':'Affect (3)','transformer_only':'Transformer scores (45)','nrc_plus_vad':'Words + affect (13)','nrc_plus_transformer':'Words + transformers (55)','vad_plus_transformer':'Affect + transformers (48)'}
for model in ['logistic_regression','xgboost','random_forest']:
 rr=[r for r in ABLATION if r['model_id']==model];full=next(r for r in rr if r['variant']=='full_hybrid');out=[full['model_name']]
 for metric in ['macro_f1','macro_average_precision']:
  best=max(rr,key=lambda r:float(r[metric]));out.append(f"{float(full[metric]):.4f} → {float(best[metric]):.4f}\n{short[best['variant']]}")
 rows.append(out)
t=table(s,['Classifier','F1: full 58 → highest observed','AP: full 58 → highest observed'],rows,y=3.15,widths=[3.2,4.41,4.42],row_h=.77,size=20).table
t.rows[0].height=Inches(.6)
for r in list(t.rows)[1:]:r.height=Inches(.82)
band(s,'The best observed feature choice depends on both the classifier and the evaluation metric.',y=6.4,size=19,h=.42)
#22
s=base(22,'Generated reviews test an additional training-data strategy','Synthetic augmentation','Thesis §3.2.2.')
steps=[('Sensory keywords\n+ templates','light'),('GPT-5\nreview text','paleorange'),('Gemini\nemotion labels','light'),('Additional\ntraining examples','navy')]
for i,(lab,fill) in enumerate(steps):
 x=.65+i*3.11;node(s,x,2.51,2.69,1.35,lab,fill,20)
 if i<3:arrow(s,x+2.75,3.18,x+3.04,3.18)
txt(s,.7,4.43,3.49,.78,'2,352',43,'orange',True);txt(s,4.23,4.44,8.3,.92,'reviews in the primary corpus\nThree generated reviews per perfume',24)
band(s,'Existing target emotion labels and VAD values were not supplied to GPT-5.',size=22)
#23
s=base(23,'Both training conditions face the same natural test data','Controlled augmentation comparison','Thesis §§3.8 and 4.1.1. Base weighting is before model-specific class weighting.')
node(s,.65,2.1,3.66,.92,'Natural training data',size=22);node(s,.65,3.6,3.66,1.13,'Natural + permitted\nsynthetic reviews','paleorange',22)
arrow(s,4.39,2.56,5.01,2.56);arrow(s,4.39,4.16,5.01,4.16)
node(s,5.1,2.1,3.05,.92,'Natural-only model',size=21);node(s,5.1,3.6,3.05,1.13,'Augmented model','paleorange',21)
arrow(s,8.24,2.58,9.05,3.13);arrow(s,8.24,4.16,9.05,3.69);node(s,9.15,2.84,3.52,1.31,'Same held-out\nnatural test perfumes','navy',22)
txt(s,.7,5.07,12,.85,'Per training perfume: natural records have total base weight 1;\nthe three synthetic reviews together also have total base weight 1.',22)
band(s,'Same folds, settings and cutoff. No synthetic reviews from validation or test perfumes.',size=20)
#24

s=base(24,'The effect of augmentation','Augmentation results','Thesis Tables 9 and 10. Pooled out-of-fold Macro AP on the same natural test data.')
txt(s,.72,1.94,11,.33,'Macro Average Precision  ·  Natural-only vs. weighted augmented training',17,'muted')
chart(s,['Logistic\nRegression','XGBoost','Random\nForest','BERT','BiLSTM +\nAttention'],[('Natural only',[.3145,.3094,.3346,.3590,.2753]),('Natural + synthetic',[.3505,.3108,.3145,.3210,.2772])],y=2.39,h=3.63,colors=['teal','orange'])
band(s,'Observed AP rises for Logistic Regression and falls for Random Forest and BERT.',y=6.27,size=20,h=.54)
#25

s=base(25,'How certain is the effect of augmentation?','From observed changes to supported conclusions','Thesis §4.4.3; Table 11. Seven-label Macro AP; 95% paired source-group intervals. Exact values: A4.')
txt(s,.69,1.95,12,.85,'Change = augmented Macro AP − natural-only Macro AP.\nPositive means improvement; negative means decline.',22)
card(s,.65,2.93,5.85,2.38,'How we obtain the range','Recalculate the difference for many samples of source groups, each given equal weight.',body_size=21)
card(s,6.83,2.93,5.85,2.38,'What the ranges show','For all five models, the range includes both negative and positive AP changes.',body_size=22)
txt(s,.73,5.54,11.9,.57,'Zero means no change. Ranges crossing zero leave the direction uncertain.',20)
band(s,'The observed gains need further confirmation; augmentation should be validated for each model.',y=6.29,size=20,h=.53)
#26

s=base(26,'What the experiments establish together','Answers to the research questions','Thesis §§6.2–6.3. Observed differences are interpreted within the tested configurations.')
items=[
 (2.03,'RQ1','Representation makes the comparison possible','58 features organise word associations, affect and contextual scores into shared inputs for three classifiers.'),
 (3.52,'RQ2','Different measures reveal different strengths','BERT leads F1 and AP; the 58-feature XGBoost route leads ROC-AUC. Label decisions and score rankings measure different aspects.'),
 (5.01,'RQ3','More training text must be tested on natural reviews','With the same natural test data, augmentation raises Logistic Regression’s AP and lowers BERT’s. A general gain is not established.')]
for y,rq,title,body in items:
 node(s,.65,y,1.12,1.18,rq,'navy',23);txt(s,2.02,y,10.5,.38,title,22,'teal',True);txt(s,2.02,y+.47,10.5,.88,body,21)
#27

s=base(27,'How far do the findings apply?','Controls and remaining boundaries','Thesis §5.6. Source isolation and sensitivity checks address specific limitations; independent validation remains necessary.')
rows=[
 ['Reference labels','Fixed Gemini prompt and shared labels','Human agreement and excerpt relevance remain unvalidated.'],
 ['Shared sources','Shared videos and identical transcripts stay together','Reviewer or channel dependence may remain.'],
 ['Rare emotions','Per-emotion results; Fear-omitted sensitivity check','Few positive cases limit the precision of rare-emotion results.'],
 ['Data domain','Evaluation on held-out English YouTube perfumes','New languages, platforms and sources need separate testing.'],
 ['Model setup','Shared test folds; fixed settings within each contrast','Input windows, pretraining and weighting differ across models.']]
t=table(s,['Boundary','Control or analysis in this study','What remains outside the claim'],rows,y=2.02,widths=[2.39,4.58,5.06],row_h=.72,size=18).table
t.rows[0].height=Inches(.58)
band(s,'Supported claim: performance of these configured approaches on this labelled review dataset.',y=6.35,size=19,h=.47)
#28

s=base(28,'What evidence would strengthen these conclusions?','Next studies follow from the boundaries','Thesis §6.4. These are proposed extensions, not completed validation results.')
txt(s,.69,1.97,12,.51,'Each next study addresses a specific boundary on the previous slide.',23,'teal',True)
rows=[
 ['Model-based labels','Check excerpt relevance; obtain independent human annotations','How closely the labels agree with human judgments'],
 ['Shared sources and\nrare-emotion cases','Collect independent sources and more examples of rare emotions','Generalisation and more precise rare-emotion estimates'],
 ['Augmentation design','Compare natural and generated data at matching units; test generation strategies','When added text improves predictions on independent natural reviews']]
t=table(s,['Current boundary','Next study','What it would establish'],rows,y=2.48,widths=[2.9,4.6,4.53],row_h=.95,size=20).table
t.rows[0].height=Inches(.57)
for r in list(t.rows)[1:]:r.height=Inches(.99)
band(s,'The current benchmark supplies the comparison framework for these next validation steps.',y=6.39,size=19,h=.43)
#29

s=base(29,'Three takeaways','Conclusion','Thesis §§6.1–6.3.',dark=True)
items=[(2.05,'01','A reproducible representation','58 named features connect lexical, affective and contextual information.'),(3.42,'02','A clear comparison of two input routes','Three feature-based and two text-based classifiers are evaluated on the same held-out perfumes.'),(4.79,'03','A measured test of additional training data','The augmentation experiment identifies model-specific changes and the limits of the supporting evidence.')]
for y,n,title,body in items:
 txt(s,.72,y,.68,.5,n,26,'blue',True);txt(s,1.63,y,10.77,.46,title,25,'white',True);txt(s,1.63,y+.57,10.77,.8,body,21,'white')
txt(s,.74,6.34,11.83,.56,'A traceable dataset, a reproducible workflow and controlled evidence.',21,'white',True)
#30
s=base(30,'Thank you','Thesis defense',dark=True)
for sh in s.shapes:
 if sh.has_text_frame and sh.text=='Thank you':
  sh.top=Inches(1.21);sh.height=Inches(1)
  for p in sh.text_frame.paragraphs:p.font.size=Pt(48)
txt(s,.72,2.57,11.8,.72,'Questions & Discussion',31,'white')
txt(s,.72,3.72,11.8,.46,'Ece ÇANGA  |  Çankaya University, Computer Engineering',22,'white')
rect(s,.72,4.58,1.0,.04,'teal')
txt(s,.72,5.02,11.8,.47,'784 perfumes   •   5 models   •   2 training conditions',23,'white',True)
txt(s,.72,5.91,11.6,.81,'Understanding emotional language in perfume reviews\nthrough reproducible modelling and evaluation.',23,'white')
#A1
s=base('A1','Full results: natural training data','Backup · full results','Thesis Tables 9 and 10. Pooled out-of-fold predictions; seven labels; cutoff 0.30.')
table(s,['Model','Macro F1','Macro AP','Macro ROC-AUC'],[['Logistic Regression','.3592','.3145','.5473'],['XGBoost','.3260','.3094','.5683'],['Random Forest','.3306','.3346','.5638'],['BERT','.4029','.3590','.5533'],['BiLSTM with Attention','.3486','.2753','.4507']],widths=[4.4,2.4,2.4,2.83],row_h=.66,size=22,highlights={(1,3),(3,1),(3,2)})
band(s,'All five models are evaluated on the same 784 held-out YouTube perfumes.',y=6.33,size=21,h=.49)
#A2
s=base('A2','Full results: augmented training data','Backup · full results','Thesis §§3.8 and 4.1.1; Tables 9 and 10. Pooled out-of-fold results; cutoff 0.30.')
table(s,['Model','Macro F1','Macro AP','Macro ROC-AUC'],[['Logistic Regression','.3737','.3505','.5372'],['XGBoost','.3359','.3108','.5982'],['Random Forest','.3247','.3145','.5731'],['BERT','.3316','.3210','.5267'],['BiLSTM with Attention','.3466','.2772','.4490']],widths=[4.4,2.4,2.4,2.83],row_h=.6,size=22,highlights={(0,1),(0,2),(1,3)})
txt(s,.72,5.94,11.9,.48,'Three synthetic reviews per eligible perfume: each has base weight ⅓; natural records total 1.',18)
txt(s,.72,6.5,11.9,.35,'Class weighting also applies. Equal base weights do not imply identical effective influence across models.',17,'muted')
#A3
s=base('A3','Source grouping and fold sizes','Backup · evaluation splits','Thesis §3.7; Figure 5; Table 8; saved fold assignments.')
table(s,['Outer fold','Perfumes','Source groups'],[['1','334','1'],['2','113','90'],['3','113','90'],['4','112','89'],['5','112','89'],['Total','784','359']],x=.65,y=2.05,w=7.34,widths=[2.44,2.45,2.45],row_h=.61,size=22,highlights={(5,0),(5,1),(5,2)})
txt(s,8.47,2.14,4.04,1.68,'Shared videos and normalized identical full transcripts connect perfumes into groups.',22)
txt(s,8.47,4.04,4.04,1.57,'Connections are transitive. Each complete group remains within one fold.',22)
band(s,'Pooled results weight perfumes equally; supplementary analyses weight source groups equally.',y=6.5,size=19,h=.32)
#A4
s=base('A4','Augmentation uncertainty under equal-source weighting','Backup · paired uncertainty','Thesis §4.4.3; Table 11. Changes are augmented minus natural-only training.')
table(s,['Model','Δ Macro AP','Unadjusted 95% interval'],[['Logistic Regression','+.1202','[−.0366, +.1402]'],['XGBoost','+.0025','[−.0205, +.0381]'],['Random Forest','−.0244','[−.1079, +.0154]'],['BERT','−.1458','[−.1726, +.0722]'],['BiLSTM with Attention','−.0044','[−.0434, +.0274]']],widths=[4.6,2.7,4.73],row_h=.56,size=21)
txt(s,.71,5.69,11.95,.48,'Seven-label Macro AP. Equal-group weighting answers a different question from pooled scores.',18)
txt(s,.71,6.22,11.95,.7,'Across seven-label and Fear-omitted analyses, 2 of 30 intervals exclude zero.\n2,000 bootstrap samples; fitted predictions held fixed; no multiple-comparison correction.',18,'muted')
#A5
s=base('A5','Annotation policy and validation boundaries','Backup · reference labels','Thesis §3.4; Appendices 2 and 3.')
card(s,.65,2.12,5.85,3.96,'Fixed annotation policy','Gemini 3.1 Pro with fixed prompts.\n\nNatural data: 1–2 excerpts, each ≤360 characters. Synthetic data: one review.\n\nWeak or neutral evidence can produce all-zero labels.',body_size=20)
card(s,6.83,2.12,5.85,3.96,'Validation boundaries','Anger excluded: 7 natural positives and 0 synthetic positives.\n\n174 of 784 records had no name anchor found; 20 inspected traces were correct.',body_size=21)
band(s,'Traceability does not establish topical relevance or agreement with human emotion judgments.',y=6.42,size=19,h=.39)
#A6
s=base('A6','Input scope and model configuration','Backup · modelling details','Thesis §§3.5.3 and 3.6.')
table(s,['Component','Input scope'],[['Lexicon and affect dimensions','Full transcript'],['Hartmann, Savani, GoEmotions, Cardiff','First 512 characters'],['BERT','Up to 96 tokens'],['BiLSTM with Attention','Up to 208 tokens']],x=.65,y=2.05,w=7.12,widths=[4.0,3.12],row_h=.69,size=20)
txt(s,8.18,2.16,4.37,1.21,'BERT\nGoEmotions initialization; up to 3 epochs.',21)
txt(s,8.18,3.73,4.37,1.21,'BiLSTM\nTrained from scratch; up to 12 epochs.',21)
txt(s,.7,5.58,12,.66,'Text probabilities are averaged at perfume level. 4,879 chunks remain after removing two within-perfume duplicates.',19)
band(s,'The comparison includes differences in input scope and training configuration.',y=6.43,size=20,h=.39)
#A7
s=base('A7','Per-emotion F1: natural training data','Backup · per-emotion results','Thesis Table 14. Perfume-level reference support; cutoff 0.30; BiLSTM uses attention.')
table(s,['Emotion','Support','LR','XGBoost','RF','BERT','BiLSTM'],[['Joy','622','.8614','.8774','.8848','.8848','.8848'],['Trust','476','.7376','.7179','.7556','.7556','.7556'],['Anticipation','366','.6380','.5831','.6294','.6365','.6365'],['Surprise','44','.1274','.1034','.0444','.0435','.1068'],['Disgust','19','.0235','.0000','.0000','.0000','.0197'],['Sadness','11','.0435','.0000','.0000','.0000','.0261'],['Fear','3','.0833','.0000','.0000','.5000','.0108']],widths=[2.4,1.4,1.61,1.72,1.61,1.62,1.67],row_h=.51,size=19)
band(s,'Rare-emotion scores can change substantially with one prediction.',y=6.42,size=21,h=.39)
#A8
s=base('A8','Selected references','Backup · references','All eight works are cited in the thesis. Titles are linked to the source pages.')
refs=[
('Nikitina & Furuoka (2023)','Words of scents: A linguistic analysis of online perfume reviews.','https://doi.org/10.1515/lingvan-2022-0135'),
('Nandwani & Verma (2021)','A review on sentiment analysis and emotion detection from text.','https://doi.org/10.1007/s13278-021-00776-6'),
('Mohammad & Turney (2013)','Crowdsourcing a Word–Emotion Association Lexicon.','https://doi.org/10.1111/j.1467-8640.2012.00460.x'),
('Mohammad (2018)','Obtaining Reliable Human Ratings of Valence, Arousal, and Dominance for 20,000 English Words.','https://aclanthology.org/P18-1017/'),
('Demszky et al. (2020)','GoEmotions: A Dataset of Fine-Grained Emotions.','https://aclanthology.org/2020.acl-main.372/'),
('Devlin et al. (2019)','BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding.','https://aclanthology.org/N19-1423/'),
('Saito & Rehmsmeier (2015)','The Precision–Recall Plot Is More Informative than the ROC Plot When Evaluating Binary Classifiers on Imbalanced Datasets.','https://doi.org/10.1371/journal.pone.0118432'),
('Niu et al. (2024)','Rethinking Emotion Annotations in the Era of Large Language Models.','https://arxiv.org/abs/2412.07906')]
for i,(authors,title,url) in enumerate(refs):
 col=i//4;row=i%4;x=.7+col*6.18;y=1.98+row*1.19
 txt(s,x,y,5.66,.3,authors,17,'teal',True)
 sh=txt(s,x,y+.37,5.66,.81,title,16)
 for p in sh.text_frame.paragraphs:
  for r in p.runs:
   r.hyperlink.address=url;r.font.color.rgb=rgb('navy');r.font.underline=False
#A9
s=base('A9','Threshold sensitivity','Backup · prediction cutoff','Thesis §4.5.2; Table 13. Macro F1 from the same YouTube-only out-of-fold probabilities.')
table(s,['Model','0.20','0.30 official','0.50','0.70'],[['Logistic Regression','.3590','.3592','.2985','.1690'],['XGBoost','.3409','.3260','.2950','.2248'],['Random Forest','.3526','.3306','.2952','.1669'],['BERT','.3926','.4029','.2899','.0000'],['BiLSTM with Attention','.3515','.3486','.1943','.0000']],widths=[4.0,1.9,2.33,1.9,1.9],row_h=.6,size=21,highlights={(i,2) for i in range(5)})
txt(s,.71,5.96,11.9,.47,'Changing the cutoff changes the labels without retraining. AP and ROC-AUC are unchanged.',19)
txt(s,.71,6.51,11.9,.31,'The official 0.30 cutoff was specified before this sensitivity check.',19,'teal',True)
#A10
s=base('A10','Model-comparison uncertainty','Backup · evidence strength','Thesis §4.4.2. Unadjusted 95% source-group bootstrap intervals; conditional on fitted models.')
txt(s,.72,1.99,12,.43,'Supplementary analysis: equal weight for each source group',23,'teal',True)
table(s,['BERT point estimate','Macro F1','Macro AP','Macro ROC-AUC'],[['Highest observed','.4839','.4665','.5932']],y=2.72,widths=[4.08,2.4,2.4,3.15],row_h=.64,size=22)
card(s,.65,4.31,5.85,1.66,'F1 comparisons','All 10 paired intervals include zero.',body_size=21)
card(s,6.83,4.31,5.85,1.66,'Ranking comparisons','7 of 20 intervals exclude zero, mainly involving BiLSTM.',body_size=19)
band(s,'These results do not establish one model as clearly best across all measures.',y=6.35,size=21,h=.46)

prs.slides[8].notes_slide.notes_text_frame.text += "\nFeature generators: j-hartmann/emotion-english-distilroberta-base; bhadresh-savani/distilbert-base-uncased-emotion; monologg/bert-base-cased-goemotions-original; cardiffnlp/twitter-roberta-base-emotion. Gemini 3.1 Pro assigns references; GPT-5 generates synthetic text. Numerical evaluation compares saved predictions with reference labels."
prs.slides[10].notes_slide.notes_text_frame.text += "\nExact saved feature columns: " + "; ".join(FEATURE_COLUMNS) + "\nThe input schema includes Anger, neutral and other categories; the supervised benchmark still predicts seven target emotions. Repeated names from different sources are distinct features."
prs.slides[24].notes_slide.notes_text_frame.text += "\nThe equal-source weighting differs from the pooled chart on slide 24. An interval crossing zero does not prove equality. The bootstrap is conditional on saved fitted-model predictions and intervals are unadjusted for multiple comparisons. The mechanism behind model-specific responses was not isolated."
# Store sources and an accessible plain-text extraction for internal validation.
assert len(prs.slides)==40
for i,s in enumerate(prs.slides):
 manifest[i]['text']=[sh.text for sh in s.shapes if sh.has_text_frame]
 manifest[i]['tables']=[[[c.text for c in row.cells] for row in sh.table.rows] for sh in s.shapes if sh.has_table]
# Backups remain available after slide 30; the closing slide links directly to them.
# Navigation is transparent and does not change the 30-slide main narrative.
nav=txt(prs.slides[29],11.25,6.79,1.41,.16,'BACKUP SLIDES →',9,'blue',align=PP_ALIGN.RIGHT)
nav.click_action.target_slide=prs.slides[30]
prs.save(PPTX)
(OUT/'build'/'revised_slide_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
print(PPTX)
print(f'{len(prs.slides)} slides; {sum(sh.has_chart for s in prs.slides for sh in s.shapes)} editable charts')
